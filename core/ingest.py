"""Ingestion pipeline: document -> units of text -> chunks -> embeddings -> ChromaDB.

A "unit" is what citations point to: a page (PDF), a slide (PowerPoint) or a part
(Word, split at headings, since .docx files have no fixed pages).
One Chroma collection per uploaded document. The collection name is derived from
the file's SHA-1, so re-uploading the same file reuses the existing index.
"""
from __future__ import annotations

import hashlib
import io
import re
import threading
import unicodedata
from dataclasses import dataclass, field
from functools import lru_cache
from typing import Callable

import pymupdf as fitz

from core import config

# Separators tried in order, coarsest first. Includes Arabic question mark / comma.
SEPARATORS = ["\n\n", "\n", ". ", "? ", "! ", "؟ ", "، ", "; ", " ", ""]
MIN_CHUNK_CHARS = 20
EMBED_BATCH = 32

DOCX_PART_CHARS = 3000   # a Word "part" also ends after this many characters

ProgressCb = Callable[[float, str], None]

# file extension -> citation unit
UNITS = {".pdf": "page", ".pptx": "slide", ".docx": "part"}
SUPPORTED_TYPES = [ext.lstrip(".") for ext in UNITS]


def fmt_ref(unit: str, n: int) -> str:
    """Short citation label: 'p.3', 'slide 3', 'part 3'."""
    return f"p.{n}" if unit == "page" else f"{unit} {n}"


def unit_of(filename: str) -> str:
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in UNITS:
        raise ValueError(f"Unsupported file type '{ext or filename}'. Supported: {', '.join(UNITS)}")
    return UNITS[ext]


@dataclass
class IngestResult:
    doc_id: str
    filename: str
    n_pages: int
    n_chunks: int
    empty_pages: list[int] = field(default_factory=list)
    already_indexed: bool = False
    unit: str = "page"

    @property
    def warning(self) -> str | None:
        if self.n_chunks == 0:
            if self.unit == "page":
                return (f"No extractable text in '{self.filename}'. It is probably a scanned/image-only "
                        "PDF. Run OCR on it first (e.g. `ocrmypdf in.pdf out.pdf`) and upload again.")
            return f"No text found in '{self.filename}' (only images?). Nothing was indexed."
        if self.unit == "page" and self.empty_pages and len(self.empty_pages) >= max(1, self.n_pages // 2):
            return (f"{len(self.empty_pages)}/{self.n_pages} pages of '{self.filename}' have no text "
                    "(scanned images?). Only the text pages were indexed.")
        return None


# --------------------------------------------------------------------------- chunking
def _normalize(text: str) -> str:
    text = text.replace("\r", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _split_recursive(text: str, seps: list[str], size: int) -> list[str]:
    if len(text) <= size:
        return [text]
    sep = next(s for s in seps if s == "" or s in text)
    if sep == "":
        return [text[i:i + size] for i in range(0, len(text), size)]
    raw = text.split(sep)
    # keep the separator attached so merged chunks read naturally
    parts = [p + sep for p in raw[:-1]] + [raw[-1]]
    rest = seps[seps.index(sep) + 1:]
    out: list[str] = []
    for p in parts:
        if not p:
            continue
        out.extend([p] if len(p) <= size else _split_recursive(p, rest, size))
    return out


def _merge(pieces: list[str], size: int, overlap: int) -> list[str]:
    chunks: list[str] = []
    cur: list[str] = []
    cur_len = 0
    for p in pieces:
        if cur and cur_len + len(p) > size:
            chunks.append("".join(cur).strip())
            # drop pieces from the front until only ~overlap chars remain to carry over
            while cur and (cur_len > overlap or cur_len + len(p) > size):
                cur_len -= len(cur.pop(0))
        cur.append(p)
        cur_len += len(p)
    if cur:
        chunks.append("".join(cur).strip())
    return [c for c in chunks if len(c) >= MIN_CHUNK_CHARS]


def split_text(text: str, chunk_size: int = config.CHUNK_SIZE,
               overlap: int = config.CHUNK_OVERLAP) -> list[str]:
    """Recursive character splitter (~chunk_size chars, ~overlap chars shared between neighbours)."""
    text = _normalize(text)
    if not text:
        return []
    return _merge(_split_recursive(text, SEPARATORS, chunk_size), chunk_size, overlap)


# --------------------------------------------------------------------------- PDF parsing
def extract_pages(pdf_bytes: bytes) -> list[tuple[int, str]]:
    """Return [(page_number starting at 1, text)] for every page.

    NFKC folds Arabic presentation forms (U+FBxx/U+FExx, which many PDFs emit for shaped
    glyphs) back to normal letters, and expands ligatures like 'ﬁ'. Without it Arabic
    text neither embeds nor matches queries properly."""
    with fitz.open(stream=pdf_bytes, filetype="pdf") as doc:
        return [(i + 1, unicodedata.normalize("NFKC", page.get_text("text"))) for i, page in enumerate(doc)]


def _shape_text(shape) -> list[str]:
    """Text from one PowerPoint shape, including tables and grouped shapes."""
    out = []
    if getattr(shape, "shape_type", None) == 6 or hasattr(shape, "shapes"):  # 6 = group
        for s in getattr(shape, "shapes", []):
            out += _shape_text(s)
    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
        out.append(shape.text_frame.text.strip())
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                out.append(" | ".join(cells))
    return out


def extract_slides(pptx_bytes: bytes) -> list[tuple[int, str]]:
    """[(slide_number, text)]: slide text, tables, then the speaker notes (often the richest part)."""
    from pptx import Presentation
    prs = Presentation(io.BytesIO(pptx_bytes))
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        parts = [t for shape in slide.shapes for t in _shape_text(shape)]
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                parts.append("Speaker notes: " + notes)
        slides.append((i, unicodedata.normalize("NFKC", "\n\n".join(parts))))
    return slides


def extract_docx_parts(docx_bytes: bytes) -> list[tuple[int, str]]:
    """[(part_number, text)]. A new part starts at each heading, or after ~3000 characters."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = Document(io.BytesIO(docx_bytes))
    parts: list[str] = []
    cur: list[str] = []

    def flush():
        if "".join(cur).strip():
            parts.append("\n\n".join(cur))
        cur.clear()

    for el in doc.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, doc)
            text = p.text.strip()
            if not text:
                continue
            if (p.style is not None and p.style.name.lower().startswith(("heading", "title"))) \
                    or sum(map(len, cur)) > DOCX_PART_CHARS:
                flush()
            cur.append(text)
        elif tag == "tbl":
            for row in Table(el, doc).rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    cur.append(" | ".join(dict.fromkeys(cells)))  # merged cells repeat their text
    flush()
    return [(i, unicodedata.normalize("NFKC", t)) for i, t in enumerate(parts, start=1)]


def extract_units(file_bytes: bytes, filename: str) -> tuple[str, list[tuple[int, str]]]:
    unit = unit_of(filename)
    reader = {"page": extract_pages, "slide": extract_slides, "part": extract_docx_parts}[unit]
    return unit, reader(file_bytes)


def chunk_pages(pages: list[tuple[int, str]]) -> tuple[list[dict], list[int]]:
    """Chunk page by page so every chunk maps to exactly one page number."""
    chunks, empty = [], []
    for page_no, text in pages:
        if len(text.strip()) < config.MIN_PAGE_CHARS:
            empty.append(page_no)
            continue
        for i, c in enumerate(split_text(text)):
            chunks.append({"text": c, "page": page_no, "chunk_index": i})
    return chunks, empty


# --------------------------------------------------------------------------- embeddings / store
_embedder_lock = threading.Lock()


@lru_cache(maxsize=1)
def _load_embedder():
    if config.MODEL_BACKEND == "onnx":
        from core.onnx_backend import OnnxEmbedder
        return OnnxEmbedder(config.EMBED_MODEL)
    from sentence_transformers import SentenceTransformer
    return SentenceTransformer(config.EMBED_MODEL, device="cpu")


def get_embedder():
    with _embedder_lock:  # avoid loading the model twice when Streamlit threads race
        return _load_embedder()


def embed_passages(texts: list[str]) -> list[list[float]]:
    # e5 models are trained with "passage: " / "query: " prefixes.
    vecs = get_embedder().encode([f"passage: {t}" for t in texts], normalize_embeddings=True,
                                 batch_size=EMBED_BATCH)
    return vecs.tolist()


def embed_query(text: str) -> list[float]:
    return get_embedder().encode(f"query: {text}", normalize_embeddings=True).tolist()


_client_lock = threading.Lock()
_clients: dict[str, object] = {}


def get_client(path: str = config.CHROMA_DIR):
    # Streamlit can run two script threads at once (a rerun while the first run is still going).
    # Building two PersistentClients concurrently yields a half-started client, hence the lock.
    with _client_lock:
        if path not in _clients:
            import chromadb
            from chromadb.config import Settings
            _clients[path] = chromadb.PersistentClient(path=path, settings=Settings(anonymized_telemetry=False))
        return _clients[path]


def doc_id_for(pdf_bytes: bytes) -> str:
    return "doc_" + hashlib.sha1(pdf_bytes).hexdigest()[:16]


def ingest_file(file_bytes: bytes, filename: str, progress: ProgressCb | None = None,
                client=None) -> IngestResult:
    """Index a PDF, PowerPoint or Word file."""
    progress = progress or (lambda f, m: None)
    client = client or get_client()
    doc_id = doc_id_for(file_bytes)

    progress(0.05, f"Reading {filename}")
    unit, pages = extract_units(file_bytes, filename)
    chunks, empty = chunk_pages(pages)
    result = IngestResult(doc_id, filename, len(pages), len(chunks), empty, unit=unit)

    existing = {getattr(c, "name", c) for c in client.list_collections()}
    if doc_id in existing and client.get_collection(doc_id).count() > 0:
        result.already_indexed = True
        progress(1.0, f"{filename} already indexed")
        return result
    if not chunks:
        progress(1.0, "No text found")
        return result

    col = client.get_or_create_collection(
        doc_id, metadata={"filename": filename, "n_pages": len(pages), "unit": unit, "hnsw:space": "cosine",
                          "embedder": config.EMBEDDER_ID})
    for start in range(0, len(chunks), EMBED_BATCH):
        batch = chunks[start:start + EMBED_BATCH]
        progress(0.1 + 0.9 * start / len(chunks), f"Embedding chunks {start + 1}-{start + len(batch)} of {len(chunks)}")
        col.add(
            ids=[f"{doc_id}-p{c['page']}-c{c['chunk_index']}" for c in batch],
            documents=[c["text"] for c in batch],
            embeddings=embed_passages([c["text"] for c in batch]),
            metadatas=[{"filename": filename, "page": c["page"], "unit": unit, "chunk_index": c["chunk_index"],
                        "doc_id": doc_id} for c in batch],
        )
    progress(1.0, f"Indexed {len(chunks)} chunks from {filename}")
    return result


def ingest_pdf(pdf_bytes: bytes, filename: str, progress: ProgressCb | None = None, client=None) -> IngestResult:
    """Backwards-compatible name used by tests and the eval script."""
    return ingest_file(pdf_bytes, filename, progress, client)


def list_documents(client=None) -> list[dict]:
    client = client or get_client()
    docs = []
    for c in client.list_collections():
        col = client.get_collection(getattr(c, "name", c))
        meta = col.metadata or {}
        docs.append({"doc_id": col.name, "filename": meta.get("filename", col.name),
                     "n_pages": meta.get("n_pages", 0), "unit": meta.get("unit", "page"), "n_chunks": col.count()})
    return sorted(docs, key=lambda d: d["filename"].lower())


def reembed_outdated(client=None, progress: ProgressCb | None = None) -> list[str]:
    """Re-embed documents that were indexed with a different embedding model or backend.

    Query and passage vectors must come from the same model; mixing (e.g. a full-precision index
    with int8 queries) quietly degrades search. Documents indexed before this check existed have no
    'embedder' entry and are treated as the old default (full precision). Returns the file names updated."""
    client = client or get_client()
    progress = progress or (lambda f, m: None)
    updated = []
    for c in client.list_collections():
        col = client.get_collection(getattr(c, "name", c))
        meta = dict(col.metadata or {})
        if meta.get("embedder", f"{config.EMBED_MODEL}|torch") == config.EMBEDDER_ID or col.count() == 0:
            continue
        name = meta.get("filename", col.name)
        data = col.get(include=["documents"])
        for start in range(0, len(data["ids"]), EMBED_BATCH):
            progress(start / len(data["ids"]), f"Updating the search index for {name}")
            col.update(ids=data["ids"][start:start + EMBED_BATCH],
                       embeddings=embed_passages(data["documents"][start:start + EMBED_BATCH]))
        meta["embedder"] = config.EMBEDDER_ID
        col.modify(metadata={k: v for k, v in meta.items() if not k.startswith("hnsw:")})
        updated.append(name)
    return updated


def delete_document(doc_id: str, client=None) -> None:
    (client or get_client()).delete_collection(doc_id)
