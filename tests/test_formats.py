import io

import pytest

from core.ingest import extract_docx_parts, extract_slides, fmt_ref, ingest_file, unit_of
from core.retriever import retrieve


def make_pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Routing basics"
    s1.placeholders[1].text = "Distance-vector routing shares tables with neighbours."
    s1.notes_slide.notes_text_frame.text = "Bellman-Ford is the algorithm behind RIP."
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Link-state routing"
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Protocol", "Algorithm"
    table.cell(1, 0).text, table.cell(1, 1).text = "OSPF", "Dijkstra shortest path"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_docx() -> bytes:
    from docx import Document
    d = Document()
    d.add_heading("Introduction to Databases", 1)
    d.add_paragraph("A relational database stores data in tables made of rows and columns.")
    d.add_heading("Normalization", 1)
    d.add_paragraph("Normalization removes redundancy. Third normal form removes transitive dependencies.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Form", "Rule"
    t.cell(1, 0).text, t.cell(1, 1).text = "1NF", "Atomic values only"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_unit_labels():
    assert (fmt_ref("page", 3), fmt_ref("slide", 4), fmt_ref("part", 2)) == ("p.3", "slide 4", "part 2")
    assert unit_of("Lecture.PPTX") == "slide" and unit_of("notes.docx") == "part"
    with pytest.raises(ValueError):
        unit_of("photo.jpg")


def test_slides_include_tables_and_speaker_notes():
    slides = extract_slides(make_pptx())
    assert [n for n, _ in slides] == [1, 2]
    assert "Speaker notes: Bellman-Ford" in slides[0][1]
    assert "OSPF | Dijkstra shortest path" in slides[1][1]


def test_docx_splits_at_headings_and_keeps_tables():
    parts = extract_docx_parts(make_docx())
    assert len(parts) == 2
    assert parts[0][1].startswith("Introduction to Databases")
    assert "1NF | Atomic values only" in parts[1][1]


def test_pptx_ingests_with_slide_citations(chroma):
    res = ingest_file(make_pptx(), "routing.pptx", client=chroma)
    assert res.unit == "slide" and res.n_pages == 2 and res.warning is None
    hit = retrieve("Which algorithm does OSPF use?", [res.doc_id], k=1, client=chroma)[0]
    assert (hit["page"], hit["unit"], hit["ref"]) == (2, "slide", "slide 2")


def test_docx_ingests_with_part_citations(chroma):
    res = ingest_file(make_docx(), "databases.docx", client=chroma)
    hit = retrieve("What does third normal form remove?", [res.doc_id], k=1, client=chroma)[0]
    assert hit["ref"] == "part 2"
