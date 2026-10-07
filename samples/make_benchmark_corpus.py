"""Build the benchmark corpus used by eval/benchmark.py.

Besides samples/networks_lecture.pdf (the transport layer), this creates:
  samples/network_layer_lecture.pdf  - Lecture 6, the network layer (11 pages, last one in Arabic)
  samples/routing_slides.pptx        - 3 slides with speaker notes

The network-layer lecture deliberately overlaps with the transport lecture (header sizes, checksums,
port numbers, Dijkstra/Bellman-Ford), so retrieval has to tell similar facts apart.
Run:  python samples/make_benchmark_corpus.py
"""
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
from samples.make_sample_pdf import build as build_transport, render  # noqa: E402

NETWORK_LAYER_PDF = HERE / "network_layer_lecture.pdf"
ROUTING_PPTX = HERE / "routing_slides.pptx"

PAGES = [
    ("Lecture 6: The Network Layer",
     "Course: CS 341 Computer Networks.\n\n"
     "The network layer moves packets from a sending host to a receiving host, across many routers. It has "
     "two key functions. Forwarding is a router-local action: when a packet arrives on an input link, the "
     "router moves it to the appropriate output link, using its forwarding table. This is the data plane and "
     "happens in nanoseconds, usually in hardware. Routing is a network-wide process that determines the "
     "end-to-end paths packets take from source to destination. This is the control plane and runs in software "
     "over seconds.\n\n"
     "The Internet's network layer offers a single service model called best effort: there is no guarantee "
     "that a datagram is delivered, no guarantee on timing or order, and no guaranteed bandwidth."),
    ("1. The IPv4 Datagram Format",
     "An IPv4 header is 20 bytes long when it carries no options. Its fields include the version number, the "
     "header length, the type of service, and the total length, a 16-bit field, so a datagram is at most 65,535 "
     "bytes. The identifier, flags and fragment offset fields support fragmentation.\n\n"
     "The time-to-live (TTL) field is decremented by one at every router; when it reaches zero the datagram is "
     "dropped. TTL ensures that datagrams do not circulate forever in a routing loop. The protocol field tells "
     "the destination which transport protocol should receive the data: 6 means TCP and 17 means UDP.\n\n"
     "The header checksum covers only the IP header, not the data, and must be recomputed at every router "
     "because the TTL changes at each hop. Finally the header holds the 32-bit source and destination IP "
     "addresses."),
    ("2. Fragmentation and the MTU",
     "The maximum amount of data a link-layer frame can carry is called the maximum transmission unit (MTU). "
     "Ethernet frames carry up to 1,500 bytes. When an IPv4 datagram is larger than the MTU of the outgoing link, "
     "a router can split it into smaller fragments.\n\n"
     "Each fragment carries the same identifier as the original datagram; the fragment offset field gives the "
     "position of the fragment's data in units of 8 bytes, and the last fragment has its more-fragments flag "
     "set to 0. Fragments are reassembled only at the destination host, never at intermediate routers.\n\n"
     "IPv6 does not allow routers to fragment. A router that receives an IPv6 datagram that is too large drops "
     "it and sends back an ICMP 'Packet Too Big' message; senders use path MTU discovery instead."),
    ("3. IPv4 Addressing and CIDR",
     "An IPv4 address is a 32-bit number, usually written in dotted-decimal notation such as 193.32.216.9. An "
     "address belongs to an interface, not to a host, so a router with several interfaces has several "
     "addresses.\n\n"
     "Classless Inter-Domain Routing (CIDR) writes a block of addresses as a.b.c.d/x, where the first x bits are "
     "the network prefix. For example 223.1.1.0/24 contains 256 addresses, of which 254 can be given to hosts. "
     "All interfaces in the same subnet share the prefix and can reach each other without a router.\n\n"
     "Three ranges are reserved for private networks: 10.0.0.0/8, 172.16.0.0/12 and 192.168.0.0/16. Blocks of "
     "public addresses are allocated by ICANN through regional registries."),
    ("4. DHCP: Getting an Address Automatically",
     "The Dynamic Host Configuration Protocol (DHCP) gives a host its IP address automatically when it joins a "
     "network, which is why it is called a plug-and-play protocol. DHCP exchanges four messages, often "
     "remembered as DORA: Discover, Offer, Request and ACK. Because the new host has no address yet, the first "
     "messages are broadcast.\n\n"
     "DHCP runs over UDP: the server listens on port 67 and the client uses port 68. Besides the IP address, "
     "the server tells the client its subnet mask, the address of its first-hop router (the default gateway) "
     "and the address of its local DNS server. Addresses are leased for a limited time and must be renewed."),
    ("5. Network Address Translation (NAT)",
     "NAT lets all the devices of a home network share a single public IP address. Inside the home, devices "
     "use private addresses such as 192.168.1.x; the NAT router replaces the private source address and port "
     "of every outgoing datagram with its own public address and a new port number, and records the mapping in "
     "a NAT translation table. Replies are translated back using the same table.\n\n"
     "Because the port field is 16 bits, one public address can support more than 60,000 simultaneous "
     "connections. NAT is controversial: routers should only process headers up to the network layer, and "
     "NAT breaks the end-to-end argument, for example for servers behind a NAT. IPv6's huge address space "
     "reduces the need for it."),
    ("6. IPv6",
     "IPv6 was designed mainly because the 32-bit IPv4 address space was running out. IPv6 addresses are 128 "
     "bits long, written as eight groups of four hexadecimal digits.\n\n"
     "The IPv6 header has a fixed length of 40 bytes, which lets routers process it faster. Several IPv4 "
     "features were removed: routers no longer fragment datagrams, and there is no header checksum, because "
     "the transport and link layers already check for errors and recomputing a checksum at every hop was "
     "costly. A new flow label field identifies datagrams that belong to the same flow.\n\n"
     "Because the whole Internet cannot switch at once, IPv6 is deployed with tunneling: an IPv6 datagram is "
     "carried as the payload of an IPv4 datagram between IPv6 routers that are connected through IPv4 routers."),
    ("7. Routing Algorithms",
     "A link-state algorithm is global: every router learns the complete topology and all link costs, and then "
     "computes least-cost paths with Dijkstra's algorithm. With n routers the basic algorithm needs O(n^2) "
     "operations. OSPF is a link-state protocol.\n\n"
     "A distance-vector algorithm is decentralized and iterative: each router only knows the costs to its "
     "neighbours and the distance vectors they send. It uses the Bellman-Ford equation "
     "dx(y) = min over neighbours v of { c(x,v) + dv(y) }. RIP is a distance-vector protocol.\n\n"
     "Distance-vector routing reacts quickly to good news but slowly to bad news: after a link cost increases, "
     "routers can keep raising their estimates step by step, which is called the count-to-infinity problem. "
     "Poisoned reverse avoids loops between two neighbours by advertising an infinite distance."),
    ("8. OSPF and BGP",
     "The Internet is a network of autonomous systems (ASes), each run by one organization such as an ISP. "
     "Routing inside an AS is intra-AS routing; OSPF is the most common choice. OSPF floods link-state "
     "advertisements to all routers in the AS, supports authentication, and can divide a large AS into areas.\n\n"
     "Routing between ASes uses the Border Gateway Protocol (BGP), the glue that holds the Internet together. "
     "BGP is a path-vector protocol: a route advertisement contains a prefix plus attributes such as AS-PATH, "
     "the list of ASes the advertisement passed through, and NEXT-HOP. eBGP sessions connect routers in "
     "different ASes, and iBGP spreads the routes inside an AS. BGP sessions run over TCP on port 179. With "
     "hot-potato routing, a router picks the gateway with the lowest intra-AS cost."),
    ("9. ICMP, ping and traceroute",
     "The Internet Control Message Protocol (ICMP) is used by hosts and routers to report errors and exchange "
     "network-layer information. ICMP messages are carried inside IP datagrams with protocol number 1. Each "
     "message has a type and a code: type 3 means destination unreachable and type 11 means the TTL expired.\n\n"
     "The ping program sends an ICMP echo request (type 8) and the target answers with an echo reply (type 0).\n\n"
     "Traceroute discovers the routers on the path to a destination. It sends UDP segments with TTL values of "
     "1, 2, 3 and so on; the n-th router discards the datagram whose TTL reaches zero and returns an ICMP "
     "'TTL expired' message, which reveals its address and the round-trip time. The destination finally "
     "answers with a 'port unreachable' message."),
]

ARABIC_TITLE = "ملخص المحاضرة: طبقة الشبكة"
ARABIC_BODY = (
    "تنقل طبقة الشبكة الحزم من جهاز المرسل إلى جهاز المستقبل عبر الموجهات. وظيفتها الأولى هي "
    "التمرير (Forwarding) داخل الموجه الواحد، ووظيفتها الثانية هي التوجيه (Routing) لتحديد المسار الكامل.<br/><br/>"
    "رأس IPv4 طوله عشرون بايتاً بدون خيارات، بينما رأس IPv6 ثابت الطول وهو أربعون بايتاً ولا يحتوي على "
    "مجموع اختباري. يمنع حقل TTL بقاء الحزمة في الشبكة إلى الأبد.<br/><br/>"
    "يسمح NAT لعدة أجهزة في المنزل بمشاركة عنوان IP عام واحد، ويمنح بروتوكول DHCP الجهاز عنواناً تلقائياً "
    "عند اتصاله بالشبكة."
)

SLIDES = [
    ("Routing basics",
     "Distance-vector routing shares tables with neighbours.",
     "Bellman-Ford is the algorithm behind RIP. Routers exchange distance vectors every 30 seconds."),
    ("Link-state routing",
     "Each router floods link-state advertisements and computes shortest paths.",
     "OSPF uses Dijkstra's shortest path algorithm. Link-state databases must be identical on all routers."),
    ("Comparing the two",
     "Link state: global knowledge, fast convergence. Distance vector: local knowledge, may count to infinity.",
     "Exam tip: know which protocol uses which algorithm, and why distance vector can loop."),
]


def build_network_layer(out: Path = NETWORK_LAYER_PDF) -> Path:
    return render(PAGES, ARABIC_TITLE, ARABIC_BODY, out)


def build_routing_pptx(out: Path = ROUTING_PPTX) -> Path:
    from pptx import Presentation
    prs = Presentation()
    for title, body, notes in SLIDES:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        slide.notes_slide.notes_text_frame.text = notes
    prs.save(out)
    return out


def corpus() -> list[Path]:
    """All benchmark files, built if missing."""
    from samples.make_sample_pdf import OUT as TRANSPORT_PDF
    for path, builder in [(TRANSPORT_PDF, build_transport), (NETWORK_LAYER_PDF, build_network_layer),
                          (ROUTING_PPTX, build_routing_pptx)]:
        if not path.exists():
            builder(path)
    return [TRANSPORT_PDF, NETWORK_LAYER_PDF, ROUTING_PPTX]


if __name__ == "__main__":
    for p in [build_network_layer(), build_routing_pptx()]:
        print("wrote", p)
